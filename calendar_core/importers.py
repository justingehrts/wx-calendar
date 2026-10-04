"""Import events from the legacy Content_Idea_Calendar.xlsx and from CSV."""
import calendar, csv, datetime as dt, io, json, re
from . import db

def parse_workbook(path_or_file, year):
    """Return {date: [event strings]} from the legacy workbook. One sheet per month,
    day grid in columns D:J (Sun-Sat), event text in the row below each date row
    (rows 5,7,...,15). Neighbor-month cells are de-duplicated by (date,text)."""
    from openpyxl import load_workbook
    wb = load_workbook(path_or_file, data_only=True)
    found = {}
    for m in range(1, 13):
        if calendar.month_name[m] not in wb.sheetnames: continue
        ws = wb[calendar.month_name[m]]
        first = dt.date(year, m, 1)
        start = first - dt.timedelta(days=(first.weekday() + 1) % 7)
        for w in range(6):
            for ci, col in enumerate("DEFGHIJ"):
                v = ws[f"{col}{5 + 2 * w}"].value
                if not v or not isinstance(v, str): continue
                d = start + dt.timedelta(days=7 * w + ci)
                for line in re.split(r"\n|\s{2,}", v):
                    line = line.strip()
                    if not line or line == "NOTES:" or line.lower() == "full moon": continue
                    found.setdefault(d, [])
                    if line not in found[d]: found[d].append(line)
    return found

def collapse(found):
    """Merge same-title entries on consecutive days into (title, start, end) spans."""
    by_title = {}
    for d, titles in found.items():
        for t in titles: by_title.setdefault(t, []).append(d)
    spans = []
    for t, days in by_title.items():
        days.sort(); s = e = days[0]
        for d in days[1:] + [None]:
            if d is not None and d == e + dt.timedelta(days=1): e = d; continue
            spans.append((t, s, e)); s = e = d
    return sorted(spans, key=lambda x: (x[1], x[0]))

AUTO_RE = re.compile(r"\bAvg\b|Earliest|Latest|sunrise|sunset|equilux|solstice|equinox|first freeze|last freeze|DST|begins|peak", re.I)
def could_autogenerate(title): return bool(AUTO_RE.search(title))

def import_workbook(con, path_or_file, year):
    """Insert spans, skipping ones already present. Returns number added."""
    existing = {(e["title"], e["start_date"], e["end_date"] or e["start_date"]) for e in db.events(con)}
    n = 0
    for t, s, e in collapse(parse_workbook(path_or_file, year)):
        if (t, s.isoformat(), e.isoformat()) in existing: continue
        db.add_event(con, t, s, e, flag="could be auto-generated" if could_autogenerate(t) else None); n += 1
    return n

CSV_COLS = ["start_date", "end_date", "title", "category", "time", "location", "notes", "url", "detail_only", "rule_json"]

def export_csv(con):
    out = io.StringIO(); w = csv.DictWriter(out, CSV_COLS, extrasaction="ignore"); w.writeheader()
    for e in db.events(con): w.writerow({**e, "rule_json": json.dumps(e["rule"]) if e["rule"] else ""})
    return out.getvalue()

def import_csv(con, text):
    n = 0
    for r in csv.DictReader(io.StringIO(text)):
        if not r.get("title") or not r.get("start_date"): continue
        cid = db.category_id_by_name(con, r.get("category") or "") or db.categorize(con, r["title"])
        db.add_event(con, r["title"].strip(), r["start_date"], r.get("end_date") or None, cid,
                     rule=json.loads(r["rule_json"]) if r.get("rule_json") else None,
                     time=r.get("time"), location=r.get("location"), notes=r.get("notes"), url=r.get("url"),
                     detail_only=(r.get("detail_only") or "0") not in ("0", "", "False"))
        n += 1
    return n
