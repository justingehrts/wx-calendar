import datetime as dt, calendar, re, warnings, math
warnings.filterwarnings('ignore')
from openpyxl import load_workbook
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib.colors import HexColor, white

YEAR=2026
CATS={
 'sport':('Sports',HexColor('#1F5FA8')),
 'astro':('Sky / Astronomy',HexColor('#6B3FA0')),
 'climo':('Climo records',HexColor('#D9622B')),
 'comm':('Community events',HexColor('#2E8B57')),
 'hol':('Holidays',HexColor('#555555')),
}
def cat(t):
    if re.search(r"Halloween|Veterans|Thanksgiving|Christmas|Hanukkah|DST|Labor Day|Mother|Father|Memorial Day|Juneteenth|Fourth",t): return 'hol'
    if re.search(r"Avg|Earliest|Latest|snow|freeze",t): return 'climo'
    if re.search(r"moon|peak|equilux|begins|Last sun|sunset|eclipse",t): return 'astro'
    if re.search(r"Clippers|Crew|OSU|Aviators|Tourn|Tourney|Open",t): return 'sport'
    return 'comm'

# ---- parse events from user's workbook
wb=load_workbook('Content_Idea_Calendar.xlsx')
events={}
for m in range(1,13):
    ws=wb[calendar.month_name[m]]
    first=dt.date(YEAR,m,1)
    start=first-dt.timedelta(days=(first.weekday()+1)%7)
    for w in range(6):
        er=5+2*w
        for ci,col in enumerate('DEFGHIJ'):
            v=ws[f'{col}{er}'].value
            if not v or not isinstance(v,str): continue
            d=start+dt.timedelta(days=7*w+ci)
            for line in re.split(r'\n|\s{2,}',v):
                line=line.strip()
                if line and line!='NOTES:': events.setdefault(d,[])
                if line and line!='NOTES:' and line not in events[d]: events[d].append(line)

import ephem
from zoneinfo import ZoneInfo
ET=ZoneInfo('America/New_York'); UTC=ZoneInfo('UTC')
def to_et(e):  # ephem.Date (UTC) -> aware Eastern datetime
    return e.datetime().replace(tzinfo=UTC).astimezone(ET)
OBS=ephem.Observer(); OBS.lat='39.9612'; OBS.lon='-82.9988'; OBS.elevation=235
OBS.pressure=0; OBS.horizon='-0:34'   # standard refraction; use_center=False -> upper limb (USNO definition)
import json, os
USNO=json.load(open('astro_usno.json')) if os.path.exists('astro_usno.json') else {}
def suntimes(d):
    u=USNO.get(d.isoformat())
    if u: return u['rise'].lstrip('0'),u['set'].lstrip('0')
    return suntimes_ephem(d)
def suntimes_ephem(d):
    OBS.date=ephem.Date(dt.datetime.combine(d,dt.time(0,0),tzinfo=ET).astimezone(UTC).replace(tzinfo=None))
    sr=to_et(OBS.next_rising(ephem.Sun())); ss=to_et(OBS.next_setting(ephem.Sun()))
    f=lambda x:x.strftime('%I:%M').lstrip('0')
    return f(sr),f(ss)
PRINC={}
for fn,lab in ((ephem.next_new_moon,'New'),(ephem.next_first_quarter_moon,'1st Qtr'),(ephem.next_full_moon,'Full'),(ephem.next_last_quarter_moon,'Last Qtr')):
    e=ephem.Date('2026/08/20')
    while e<ephem.Date('2027/02/01'):
        e=fn(e); PRINC[to_et(e).date()]=lab
def mphase(d):  # 0-28 scale (fraction of synodic month * 28)
    noon=dt.datetime.combine(d,dt.time(12),tzinfo=ET).astimezone(UTC).replace(tzinfo=None)
    e=ephem.Date(noon); prev=ephem.previous_new_moon(e); nxt=ephem.next_new_moon(e)
    return (e-prev)/(nxt-prev)*28
def principal(d): return PRINC.get(d)

def moon_icon(c,x,y,r,p):
    # x,y centre; p 0-28. lit fraction drawn as terminator ellipse
    c.saveState()
    c.setLineWidth(.5); c.setStrokeColor(HexColor('#777777'))
    c.setFillColor(HexColor('#2b2b2b')); c.circle(x,y,r,fill=1,stroke=1)
    ang=p/28*2*math.pi
    k=math.cos(ang)  # 1 new .. -1 full
    illum=(1-k)/2
    if illum>0.02:
        c.setFillColor(HexColor('#F4E9A8'))
        pth=c.beginPath()
        n=40
        waxing=p<14
        side=1 if waxing else -1   # lit on right when waxing
        pts=[]
        for i in range(n+1):
            t=-math.pi/2+math.pi*i/n
            pts.append((x+side*r*math.cos(t), y+r*math.sin(t)))
        # terminator
        for i in range(n+1):
            t=math.pi/2-math.pi*i/n
            pts.append((x+side*r*k*math.cos(t)*-1*-1*(-1) if False else x+side*(r*k)*math.cos(t)*1, y+r*math.sin(t)))
        pth.moveTo(*pts[0])
        for q in pts[1:]: pth.lineTo(*q)
        pth.close(); c.drawPath(pth,fill=1,stroke=0)
    c.restoreState()

def fit(c,text,font,size,maxw):
    while c.stringWidth(text,font,size)>maxw and size>4.6: size-=.2
    return size

def mini(c,x,y,w,month,year):
    c.setFont('Helvetica-Bold',7); c.setFillColor(HexColor('#222'))
    c.drawCentredString(x+w/2,y,f'{calendar.month_abbr[month].upper()} {year}')
    cw=w/7; c.setFont('Helvetica',5.5); c.setFillColor(HexColor('#888'))
    for i,l in enumerate('SMTWTFS'): c.drawCentredString(x+cw*(i+.5),y-8,l)
    cal=calendar.Calendar(6).monthdayscalendar(year,month)
    c.setFillColor(HexColor('#222'))
    for r,wk in enumerate(cal):
        for i,dn in enumerate(wk):
            if dn: c.drawCentredString(x+cw*(i+.5),y-16-r*7.2,str(dn))

import csv, os
CLIMO={}
if os.path.exists('climo.csv'):
    for r in csv.DictReader(open('climo.csv')):
        CLIMO[dt.date.fromisoformat(r['date'])]=r
def climo_line(d):
    r=CLIMO.get(d)
    if not r: return 'Nrm --/--   Rec --/--'
    t=f"Nrm {r['normal_high']}/{r['normal_low']}"
    if r.get('record_high'): t+=f"   Rec {r['record_high']}/{r['record_low']}"
    return t

def layout_week(wk,month):
    """collapse consecutive-day repeats into spanning bars; return lanes of segments"""
    names={}
    for i,d in enumerate(wk):
        for e in events.get(d,[]):
            if e.lower()=='full moon': continue
            names.setdefault(e,[]).append(i)
    segs=[]
    for e,cols in names.items():
        run=[cols[0]]
        for cidx in cols[1:]+[None]:
            if cidx is not None and cidx==run[-1]+1: run.append(cidx); continue
            segs.append((run[0],run[-1],e)); run=[cidx]
    segs.sort(key=lambda t:(t[0],-(t[1]-t[0]),t[2]))
    lanes=[]
    for sg in segs:
        for ln in lanes:
            if all(sg[0]>o[1] or sg[1]<o[0] for o in ln): ln.append(sg); break
        else: lanes.append([sg])
    return lanes

def page(c,month):
    W,H=landscape(letter); M=22
    c.setFillColor(white); c.rect(0,0,W,H,fill=1,stroke=0)
    top=H-M
    c.setFillColor(HexColor('#14213D')); c.setFont('Helvetica-Bold',34)
    c.drawString(M,top-30,calendar.month_name[month].upper())
    c.setFillColor(HexColor('#888')); c.setFont('Helvetica',15)
    c.drawString(M+c.stringWidth(calendar.month_name[month].upper(),'Helvetica-Bold',34)+8,top-30,str(YEAR))
    c.setFont('Helvetica',7); c.setFillColor(HexColor('#666'))
    c.drawString(M+2,top-42,'Weathercast planning calendar  |  Columbus, OH  |  sunrise/sunset & moon in Eastern Time  |  Nrm = normal high/low, Rec = record high/low')
    lx=M+2; ly=top-54
    for k,(n,col) in CATS.items():
        c.setFillColor(col); c.roundRect(lx,ly-1,8,8,1.5,fill=1,stroke=0)
        c.setFillColor(HexColor('#333')); c.setFont('Helvetica',7); c.drawString(lx+11,ly,n)
        lx+=c.stringWidth(n,'Helvetica',7)+26
    pm=month-1 or 12; py=YEAR-(1 if month==1 else 0)
    nm=month%12+1; ny=YEAR+(1 if month==12 else 0)
    mini(c,W-M-150,top-8,66,pm,py); mini(c,W-M-72,top-8,66,nm,ny)
    notes_h=50; hdr=14
    gtop=top-68; gbot=M+notes_h+7
    weeks=calendar.Calendar(6).monthdatescalendar(YEAR,month)
    rh=(gtop-hdr-gbot)/len(weeks); cw=(W-2*M)/7
    c.setFillColor(HexColor('#14213D')); c.rect(M,gtop-hdr,W-2*M,hdr,fill=1,stroke=0)
    c.setFillColor(white); c.setFont('Helvetica-Bold',8)
    for i,n in enumerate(['SUNDAY','MONDAY','TUESDAY','WEDNESDAY','THURSDAY','FRIDAY','SATURDAY']):
        c.drawCentredString(M+cw*(i+.5),gtop-hdr+4,n)
    for r,wk in enumerate(weeks):
        ytop=gtop-hdr-r*rh; ybot=ytop-rh
        for i,d in enumerate(wk):
            x=M+i*cw; inm=d.month==month; wkend=i in(0,6)
            c.setFillColor(HexColor('#FFFFFF') if inm and not wkend else HexColor('#F3F5F8') if inm else HexColor('#E4E6EA'))
            c.setStrokeColor(HexColor('#B8BDC6')); c.setLineWidth(.6)
            c.rect(x,ybot,cw,rh,fill=1,stroke=1)
            c.setFillColor(HexColor('#14213D') if inm else HexColor('#9AA0AA'))
            c.setFont('Helvetica-Bold',13 if inm else 10); c.drawString(x+4,ytop-14,str(d.day))
            p=mphase(d); lab=principal(d)
            moon_icon(c,x+cw-9,ytop-10,5.2,p)
            if lab:
                c.setFont('Helvetica-Bold',5.3); c.setFillColor(HexColor('#6B3FA0'))
                c.drawRightString(x+cw-17,ytop-11.5,lab.upper())
            c.setFillColor(HexColor('#9a4a1c') if inm else HexColor('#A8AEB8'))
            c.setFont('Helvetica',5.4); c.drawString(x+4,ybot+10.5,climo_line(d))
            if inm:
                sr,ss=suntimes(d)
                c.setFillColor(HexColor('#8A6D00')); c.setFont('Helvetica',5.4)
                c.drawString(x+4,ybot+3,f'rise {sr}  set {ss}')
        # event bars (spans)
        lanes=layout_week(wk,month); bh=8.2; maxl=int((rh-19-22)//(bh+1))
        for li,ln in enumerate(lanes):
            if li>=maxl: break
            by=ytop-19-(li+1)*(bh+1)
            for (c0,c1,e) in ln:
                inm=wk[c0].month==month
                col=CATS[cat(e)][1] if inm else HexColor('#A8AEB8')
                bx=M+c0*cw+3; bw=(c1-c0+1)*cw-6
                c.setFillColor(col); c.roundRect(bx,by,bw,bh,1.5,fill=1,stroke=0)
                size=fit(c,e,'Helvetica-Bold',6.6,bw-4)
                c.setFillColor(white); c.setFont('Helvetica-Bold',size); c.drawString(bx+2.5,by+2.3,e)
        if len(lanes)>maxl:
            c.setFillColor(HexColor('#333')); c.setFont('Helvetica',5.5)
            c.drawString(M+3,ytop-19-maxl*(bh+1)-6,f'+{len(lanes)-maxl} more lane(s) not shown')
    bw=(W-2*M-8)/2
    for i,t in enumerate(['CONTENT IDEAS & SEGMENTS','NOTES / TO DO']):
        x=M+i*(bw+8)
        c.setFillColor(HexColor('#E9EDF3')); c.rect(x,M,bw,notes_h,fill=1,stroke=0)
        c.setStrokeColor(HexColor('#B8BDC6')); c.setLineWidth(.6); c.rect(x,M,bw,notes_h,fill=0,stroke=1)
        c.setFillColor(HexColor('#14213D')); c.setFont('Helvetica-Bold',7.5); c.drawString(x+5,M+notes_h-9.5,t)
        c.setStrokeColor(HexColor('#C6CBD3')); c.setLineWidth(.4)
        yy=M+notes_h-21
        while yy>M+4: c.line(x+5,yy,x+bw-5,yy); yy-=10.5

c=canvas.Canvas('Weathercast_Planning_Calendar.pdf',pagesize=landscape(letter))
c.setTitle('Weathercast Planning Calendar - prototype')
for m in (10,11,12):
    page(c,m); c.showPage()
c.save()
