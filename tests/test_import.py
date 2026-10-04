import datetime as dt
from openpyxl import Workbook
from calendar_core import db, importers

def make_wb(path):
    wb = Workbook(); wb.remove(wb.active)
    for m in range(1, 13):
        wb.create_sheet(__import__("calendar").month_name[m])
    ws = wb["October"]   # Oct 1 2026 is Thursday -> grid starts Sun Sep 27
    ws["G5"] = "Clippers"; ws["H5"] = "Clippers\nHalloween"   # Sep 30? col G=Wed Sep30, H=Thu Oct1
    ws["I5"] = "Clippers  Full moon"; ws["D5"] = "NOTES:"
    wb["September"]["G13"] = "Dummy"  # noise on another sheet
    wb.save(path)

def test_import_collapse_dedupe_categorize(tmp_path):
    p = tmp_path / "x.xlsx"; make_wb(p)
    con = db.connect(str(tmp_path / "t.db"))
    assert importers.import_workbook(con, str(p), 2026) >= 2
    evs = {e["title"]: e for e in db.events(con)}
    assert evs["Clippers"]["start_date"] == "2026-09-30" and evs["Clippers"]["end_date"] == "2026-10-02"
    assert evs["Clippers"]["category"] == "Sports" and evs["Halloween"]["category"] == "Holidays"
    assert "Full moon" not in evs and "NOTES:" not in evs
    assert importers.import_workbook(con, str(p), 2026) == 0   # idempotent

def test_csv_roundtrip_and_softdelete(tmp_path):
    con = db.connect(str(tmp_path / "t.db"))
    i = db.add_event(con, "Pumpkin Show", "2026-10-21", "2026-10-25")
    txt = importers.export_csv(con)
    con2 = db.connect(str(tmp_path / "u.db")); assert importers.import_csv(con2, txt) == 1
    assert [e["title"] for e in db.events_on_day(con2, dt.date(2026, 10, 23))] == ["Pumpkin Show"]
    db.delete_event(con, i); assert db.events(con) == []

def test_wal(tmp_path):
    con = db.connect(str(tmp_path / "t.db"))
    assert con.execute("PRAGMA journal_mode").fetchone()[0] == "wal"
